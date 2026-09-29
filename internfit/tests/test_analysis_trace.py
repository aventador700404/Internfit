import contextlib
import importlib.util
import io
import json
import os
from pathlib import Path
import unittest
from unittest.mock import MagicMock, patch

from docx import Document

from core.analysis_trace import CONSENT_VERSION, build_evidence_trace, build_trace_summary
from core.cv_parser import _profile_from_lines
from core.job_parser import JobPosting
from core.llm_client import LunaResult, _validated_semantic
from core.scoring import assess_fit
from core.telemetry import emit_analysis_event
from core.telemetry_store import build_storage_row


CV_LINE = "Coordinated launch priorities with internal teams and tracked outcomes."
JOB_LINE = "Align cross-functional teams around launch priorities."


def fixtures():
    candidate = _profile_from_lines(["B.B.A. Candidate, Class of 2027", "English C1", CV_LINE], "private-name.docx")
    job = JobPosting(title="Program Intern", company="Example", url="", text=JOB_LINE)
    raw = {
        "job_core_responsibility_tags": ["stakeholder"],
        "candidate_evidence": [
            {"tag": "stakeholder", "strength": "supporting", "evidence": CV_LINE},
            {"tag": "strategy", "strength": "direct", "evidence": "Invented private sentence never in the CV."},
        ],
        "matches": [{"tag": "stakeholder", "statement": "Team coordination supports this role's launch work.", "cv_evidence": CV_LINE}],
        "gaps": [{"tag": "stakeholder", "edit_type": "clarify_existing", "suggestion": "State an outcome of the launch coordination if available.", "job_source_id": "J001", "cv_source_id": "C001"}],
    }
    validation = []
    semantic = _validated_semantic(raw, candidate, job, diagnostics=validation)
    luna = LunaResult(status="used", used=True, semantic=semantic, validation=validation)
    return candidate, job, luna


def scored():
    candidate, job, luna = fixtures()
    before, after = {}, {}
    baseline = assess_fit(candidate, job, diagnostics=before)
    result = assess_fit(candidate, job, semantic=luna.semantic, diagnostics=after)
    return candidate, job, luna, baseline, result, before, after


class AnalysisTraceTests(unittest.TestCase):
    def test_validation_reports_rejected_quotes_without_retaining_them(self):
        _, _, luna = fixtures()
        self.assertEqual([item["status"] for item in luna.validation], ["accepted", "rejected", "accepted", "accepted"])
        self.assertEqual(luna.validation[1]["reason"], "cv_quote_not_found_or_too_short")
        self.assertNotIn("Invented private sentence", json.dumps(luna.validation))
        self.assertEqual(luna.semantic["candidate"]["semantic_strengths"], {"stakeholder": 1})

    def test_same_input_comparison_and_cap_diagnostics(self):
        values = scored()
        candidate, job, luna, baseline, result, before, after = values
        summary = build_trace_summary(*values)
        self.assertEqual(summary["llm_score_delta"], result.score - baseline.score)
        self.assertEqual(summary["rule_only_score"], assess_fit(candidate, job).score)
        self.assertEqual(candidate.semantic_evidence, {})
        self.assertIn("stakeholder", after["requirements"]["responsibility_tags"])
        self.assertNotIn(CV_LINE, json.dumps(summary))
        technical = JobPosting(title="Technical Intern", company="Example", url="", text="",
                               requirements={"core_checks": {"graduate_technical_degree"}})
        diagnostic = {}
        capped = assess_fit(candidate, technical, diagnostics=diagnostic)
        self.assertEqual(diagnostic["eligibility_cap"], 45)
        self.assertEqual(capped.score, min(diagnostic["score_after_domain_cap"], 45))

    def test_evidence_links_sources_and_bounds_contact_excerpts(self):
        candidate, job, luna, baseline, result, before, after = scored()
        candidate.evidence["research"] = ["Researched a market. Contact test.person@example.com +82 10-1234-5678 https://example.com/profile " + "x" * 500]
        after["candidate_tag_strengths"]["research"] = 2
        trace = build_evidence_trace(candidate, job, luna, result, after)
        serialized = json.dumps(trace)
        self.assertNotIn("test.person@example.com", serialized)
        self.assertNotIn("10-1234-5678", serialized)
        self.assertNotIn("https://example.com/profile", serialized)
        self.assertNotIn(candidate.source_name, serialized)
        self.assertNotIn("Invented private sentence", serialized)
        references = {item["id"]: item for item in trace["excerpts"]}
        self.assertEqual(references[trace["llm_matches"][0]["cv_excerpt"]]["text"], CV_LINE)
        self.assertIn(JOB_LINE, references[trace["llm_gaps"][0]["job_excerpt"]]["text"])
        self.assertTrue(all(len(item["text"]) <= 281 for item in trace["excerpts"]))
        for source in ("cv", "job"):
            self.assertLessEqual(sum(item["source"] == source for item in trace["excerpts"]), 24)

    def test_storage_requires_boolean_consent_and_current_version(self):
        trace = {"consent_version": CONSENT_VERSION, "excerpts": [{"text": CV_LINE}]}
        base = {"event": "analysis_completed", "analysis_id": "test", "consent_version": CONSENT_VERSION}
        for consent in (None, False, "true"):
            row = build_storage_row(dict(base, evidence_storage_consent=consent, analysis_trace=trace), private_trace=trace)
            self.assertNotIn("analysis_trace", row["payload"])
        row = build_storage_row(dict(base, evidence_storage_consent=True), private_trace=trace)
        self.assertEqual(row["payload"]["analysis_trace"], trace)
        for version in ("", "old-version"):
            row = build_storage_row(dict(base, consent_version=version, evidence_storage_consent=True), private_trace=trace)
            self.assertNotIn("analysis_trace", row["payload"])

    def test_private_trace_never_reaches_stdout_even_when_storage_fails(self):
        for saved in (True, False):
            output = io.StringIO()
            trace = {"consent_version": CONSENT_VERSION, "excerpts": [{"text": CV_LINE}]}
            with patch("core.telemetry.persist_analysis_event", return_value=saved) as persist, contextlib.redirect_stdout(output):
                stored = emit_analysis_event("analysis_completed", "test", score=70,
                    evidence_storage_consent=True, consent_version=CONSENT_VERSION,
                    private_trace=trace, raw_cv_text=CV_LINE)
            self.assertEqual(stored, saved)
            self.assertEqual(persist.call_args.kwargs["private_trace"], trace)
            self.assertNotIn(CV_LINE, output.getvalue())
            self.assertNotIn("analysis_trace", json.loads(output.getvalue()))


class ServerTraceTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        path = Path(__file__).resolve().parents[2] / "internfit_web" / "server.py"
        spec = importlib.util.spec_from_file_location("internfit_trace_server", path)
        cls.server = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(cls.server)

    def request(self, consent=None, version=CONSENT_VERSION, storage_error=False):
        document = Document()
        document.add_paragraph("B.B.A. Candidate, Class of 2027")
        document.add_paragraph("English C1")
        document.add_paragraph(CV_LINE)
        file = io.BytesIO()
        document.save(file)
        parts = [("cv", file.getvalue(), "private-name.docx"), ("job_text", JOB_LINE.encode(), "")]
        if consent is not None:
            parts.extend([("evidence_storage_consent", consent.encode(), ""), ("consent_version", version.encode(), "")])
        body = b""
        for name, value, filename in parts:
            disposition = f'Content-Disposition: form-data; name="{name}"'
            if filename:
                disposition += f'; filename="{filename}"'
            body += b"--trace-boundary\r\n" + disposition.encode() + b"\r\n\r\n" + value + b"\r\n"
        body += b"--trace-boundary--\r\n"
        handler = object.__new__(self.server.AppHandler)
        handler.path = "/api/analyze"
        handler.headers = {"Content-Type": "multipart/form-data; boundary=trace-boundary", "Content-Length": str(len(body))}
        handler.rfile = io.BytesIO(body)
        handler._json = MagicMock()
        _, _, luna = fixtures()
        response = MagicMock()
        response.__enter__.return_value = response
        printed = io.StringIO()
        with patch.dict(os.environ, {"SUPABASE_URL": "https://example.supabase.co", "SUPABASE_SERVICE_ROLE_KEY": "fake-key"}, clear=True), \
             patch.object(self.server, "analyze_with_luna", return_value=luna) as analyze, \
             patch("core.telemetry_store.urlopen", return_value=response, side_effect=OSError("offline") if storage_error else None) as post, \
             contextlib.redirect_stdout(printed):
            handler.do_POST()
        analyze.assert_called_once()
        self.assertEqual(handler._json.call_args.args[0], 200)
        sent = json.loads(post.call_args.args[0].data)
        self.assertNotIn(CV_LINE, printed.getvalue())
        return handler._json.call_args.args[1], sent["payload"]

    def test_opt_in_request_persists_evidence_in_same_row_with_one_model_call(self):
        response, payload = self.request("true")
        self.assertEqual(response["evidence_storage_status"], "stored")
        self.assertIn("analysis_trace", payload)
        self.assertTrue(payload["evidence_storage_consent"])
        self.assertIn("rule_only_score", payload)
        self.assertEqual(response["score"] - payload["rule_only_score"], payload["llm_score_delta"])
        self.assertEqual(response["cv_advice_source"], "llm")
        self.assertEqual(payload["llm_gap_count"], 1)
        self.assertEqual(payload["llm_gap_status"], "accepted")
        self.assertIn("State an outcome of the launch coordination if available.", response["gap_details"])

    def test_old_clients_and_unchecked_consent_never_save_evidence(self):
        for consent, version in ((None, CONSENT_VERSION), ("false", CONSENT_VERSION), ("true", "old")):
            response, payload = self.request(consent, version)
            self.assertEqual(response["evidence_storage_status"], "not_requested")
            self.assertNotIn("analysis_trace", payload)

    def test_db_failure_preserves_analysis_and_reports_not_saved(self):
        response, payload = self.request("true", storage_error=True)
        self.assertEqual(response["evidence_storage_status"], "failed")
        self.assertIn("score", response)


if __name__ == "__main__":
    unittest.main()
